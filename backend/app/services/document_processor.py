import re
import uuid
import hashlib
from pathlib import Path
from typing import List, Optional, Tuple

from langchain_core.documents import Document

from app.config import settings


class SimpleTextSplitter:
    def __init__(self, chunk_size: int = 400, chunk_overlap: int = 80, separators: List[str] = None):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.separators = separators or ["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]

    def split_text(self, text: str) -> List[str]:
        if not text:
            return []

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size

            if end >= len(text):
                chunks.append(text[start:].strip())
                break

            best_split = end
            for sep in self.separators:
                idx = text.rfind(sep, start, end)
                if idx > start:
                    best_split = idx + len(sep)
                    break

            chunk = text[start:best_split].strip()
            if chunk:
                chunks.append(chunk)

            start = best_split - self.chunk_overlap
            if start < 0:
                start = 0

        return chunks


# 中文文档常见标题：第X章/篇（1级）、第X节（2级）、一、/（一）/1.1（3级）
_HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*(\d+)$", re.IGNORECASE)
_HEADING_TEXT = re.compile(
    r"^\s*(?:第\s*[一二三四五六七八九十百零\d]+\s*([章篇节])"
    r"|\d+(?:\.\d+)+"
    r"|[一二三四五六七八九十]+\s*、"
    r"|[(（][一二三四五六七八九十\d]+[)）])"
)


def _heading_level(line: str, style_name: Optional[str] = None) -> Optional[int]:
    """返回标题层级（1~3），非标题返回 None。"""
    text = line.strip()
    if not text or len(text) > 60:
        return None
    if style_name:
        m = _HEADING_STYLE.match(style_name.strip())
        if m:
            return min(int(m.group(1)), 3)
    # 正文列表项（如 "（三）……15天。"）常以编号开头，但带句读、偏长，不算标题
    if len(text) > 30 or any(p in text for p in "。；！？"):
        return None
    m = _HEADING_TEXT.match(text)
    if m:
        kind = m.group(1)
        if kind in ("章", "篇"):
            return 1
        if kind == "节":
            return 2
        return 3
    return None


class DocumentProcessor:
    def __init__(self):
        self.text_splitter = SimpleTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
        )

    def get_file_hash(self, file_path: Path) -> str:
        hasher = hashlib.md5()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hasher.update(chunk)
        return hasher.hexdigest()

    def generate_document_id(self) -> str:
        return f"doc_{uuid.uuid4().hex[:12]}"

    def _extract_entries(self, file_path: Path) -> List[Tuple[str, Optional[int], Optional[str]]]:
        """提取为 (行文本, 页码或None, 标题样式或None) 列表，供分节使用。"""
        suffix = file_path.suffix.lower()
        entries: List[Tuple[str, Optional[int], Optional[str]]] = []

        if suffix == ".pdf":
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        entries.append((line, page_num, None))
        elif suffix in (".docx", ".doc"):
            from docx import Document as DocxDocument

            doc = DocxDocument(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    style = para.style.name if para.style is not None else None
                    entries.append((para.text, None, style))
        else:
            raise ValueError(f"不支持的文件格式: {suffix}")
        return entries

    @staticmethod
    def _segment_lines(entries) -> List[dict]:
        """按标题层级和 PDF 页边界切成段，每段带所属章节路径与页码。"""
        segments: List[dict] = []
        stack = {}  # level -> 标题文本
        buf: List[str] = []
        cur_section: Optional[str] = None
        cur_page: Optional[int] = None

        def flush():
            text = "\n".join(buf).strip()
            buf.clear()
            if text:
                segments.append({"section": cur_section, "page": cur_page, "text": text})

        for line, page, style in entries:
            if page is not None and cur_page is not None and page != cur_page:
                flush()
            if page is not None:
                cur_page = page

            level = _heading_level(line, style)
            if level:
                flush()
                for l in [l for l in list(stack) if l >= level]:
                    del stack[l]
                stack[level] = line.strip()
                cur_section = " > ".join(stack[l] for l in sorted(stack))
            else:
                buf.append(line)
        flush()
        return segments

    def build_chunks(self, file_path: Path, filename: str, document_id: str) -> List[Document]:
        """解析文件 → 章节感知分块。重建索引与上传共用此入口。"""
        entries = self._extract_entries(file_path)
        segments = self._segment_lines(entries)

        documents: List[Document] = []
        for seg in segments:
            for piece in self.text_splitter.split_text(seg["text"]):
                idx = len(documents)
                documents.append(
                    Document(
                        page_content=piece,
                        metadata={
                            "document_id": document_id,
                            "filename": filename,
                            "source": filename,
                            "chunk_id": f"{document_id}_chunk_{idx}",
                            "chunk_index": idx,
                            "section": seg["section"],
                            "page": seg["page"],
                        },
                    )
                )

        if not documents:
            raise ValueError("文档内容为空，无法处理")
        return documents

    def process_document(self, file_path: Path, filename: str) -> Tuple[str, List[Document]]:
        document_id = self.generate_document_id()
        return document_id, self.build_chunks(file_path, filename, document_id)

    def save_uploaded_file(self, file_content: bytes, filename: str) -> Path:
        safe_filename = f"{uuid.uuid4().hex[:8]}_{filename}"
        file_path = settings.DOCUMENTS_DIR / safe_filename

        with open(file_path, 'wb') as f:
            f.write(file_content)

        return file_path

    def validate_file(self, filename: str, file_size: int) -> Tuple[bool, str]:
        ext = Path(filename).suffix.lower()

        if ext not in settings.ALLOWED_EXTENSIONS:
            return False, f"不支持的文件格式: {ext}。支持的格式: {', '.join(settings.ALLOWED_EXTENSIONS)}"

        if file_size > settings.MAX_FILE_SIZE:
            return False, f"文件大小超过限制: {file_size / 1024 / 1024:.2f}MB (最大 {settings.MAX_FILE_SIZE / 1024 / 1024}MB)"

        return True, "文件验证通过"


document_processor = DocumentProcessor()
